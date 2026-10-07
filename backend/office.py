"""常用文件在网页里看（总蓝图 S1-10 网页都能看）：视频、音频交给浏览器放；PPT 在核心里列出每页的字和图。

作者 2026-09-29：「我的网页端要能看，如果能改就更好了，但是重活需要自己安装插件」。
- Word、Excel 在网页那边用随应用带的小库显示（T15 SheetJS、T16 docx-preview），这里不管
- PPT 照原样一页页翻是插件的事（「PPT 预览」：本机 PowerPoint 或 LibreOffice 转 PDF，S1-11）；核心只用 python-pptx（T17）把字和图读出来
- 「用本机软件打开」：只开这几类文件，程序、脚本一律不开（免得被别的网页骗着运行东西）
"""
from __future__ import annotations

import os
import subprocess
import sys
import threading
from pathlib import Path

VIDEO = {".mp4", ".webm", ".ogv", ".m4v", ".mov", ".avi", ".mkv", ".wmv", ".flv"}
AUDIO = {".mp3", ".wav", ".m4a", ".ogg", ".oga", ".flac", ".aac", ".opus"}
OFFICE = {".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx", ".xls": "xlsx", ".pptx": "pptx", ".ppt": "pptx"}   # 老 .ppt 只有插件看得了
OPENABLE = VIDEO | AUDIO | set(OFFICE) | {".doc", ".csv", ".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp",
                                        ".svg", ".txt", ".md", ".tex", ".bib", ".drawio"}
MAX_TEXTS = 60                       # 一页最多列几段字
_CACHE: dict[tuple[str, int], object] = {}
_LOCK = threading.Lock()


def kind(path: Path) -> str | None:
    ext = path.suffix.lower()
    if ext in VIDEO:
        return "video"
    if ext in AUDIO:
        return "audio"
    return OFFICE.get(ext)


def _open(f: Path):
    """同一份 PPT 读一次就记住（改了才重读）；最多记两份。"""
    from pptx import Presentation
    key = (str(f), f.stat().st_mtime_ns)
    with _LOCK:
        if key not in _CACHE:
            while len(_CACHE) >= 2:
                _CACHE.pop(next(iter(_CACHE)))
            _CACHE[key] = Presentation(str(f))
        return _CACHE[key]


def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _pictures(slide) -> list:
    from pptx.shapes.picture import Picture
    return [sh for sh in _walk(slide.shapes) if isinstance(sh, Picture)]


def pptx_outline(f: Path) -> dict:
    """每一页：标题、一段段字（带缩进层级）、表格、几张图、讲稿。"""
    p = _open(f)
    out = []
    for n, s in enumerate(p.slides, 1):
        title_sh = s.shapes.title
        title = title_sh.text_frame.text.strip() if title_sh is not None and title_sh.has_text_frame else ""
        tid = title_sh.shape_id if title_sh is not None else None
        texts, tables = [], []
        for sh in _walk(s.shapes):
            if sh.shape_id == tid:
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                tables.append([[c.text.strip() for c in r.cells] for r in list(sh.table.rows)[:30]])
            elif sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip() or para.text.strip()
                    if t:
                        texts.append({"t": t, "lv": para.level})
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide and s.notes_slide.notes_text_frame else ""
        out.append({"n": n, "title": title, "texts": texts[:MAX_TEXTS], "tables": tables, "pictures": len(_pictures(s)),
                    "notes": notes[:3000]})
    return {"slides": out, "ratio": round(p.slide_width / p.slide_height, 3) if p.slide_height else 1.78}


def pptx_picture(f: Path, slide: int, i: int) -> tuple[bytes, str]:
    """第 slide 页（从 1 数）的第 i 张图（从 0 数）。"""
    p = _open(f)
    if not 1 <= slide <= len(p.slides):
        raise IndexError(slide)
    img = _pictures(p.slides[slide - 1])[i].image
    return img.blob, img.content_type


def open_local(f: Path) -> None:
    """用这台电脑上的默认程序打开（Word、Excel、PowerPoint、WPS、播放器……）。只开常见的文档、表格、图片、音视频。"""
    if f.suffix.lower() not in OPENABLE:
        raise ValueError(f"「{f.suffix}」这种文件不从网页打开（程序、脚本一律不开）")
    if sys.platform.startswith("win"):
        os.startfile(str(f))                                   # noqa: S606（只开上面这几类）
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(f)])
    else:
        subprocess.Popen(["xdg-open", str(f)])
