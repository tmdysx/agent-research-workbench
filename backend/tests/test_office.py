"""网页都能看（总蓝图 S1-10）：视频音频能放（能拖进度条）、PPT 核心列每页的字和图、Word / Excel 交给网页那边的小库；用本机软件打开只开文档和音视频。
作者 2026-09-29：「我的网页端要能看，如果能改就更好了」「这个可以看的我要做自动的，不要用agent」。"""
import io

from fastapi.testclient import TestClient

import files
import office
from main import create_app


def _pptx(path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "第一页的标题"
    s.placeholders[1].text_frame.text = "要点一"
    s.placeholders[1].text_frame.add_paragraph().text = "要点二"
    s.placeholders[1].text_frame.paragraphs[1].level = 1
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(buf, "PNG")
    buf.seek(0)
    s.shapes.add_picture(buf, Inches(1), Inches(4))
    s.notes_slide.notes_text_frame.text = "讲稿在这"
    t = p.slides.add_slide(p.slide_layouts[5])
    t.shapes.title.text = "第二页有表格"
    tb = t.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tb.cell(0, 0).text, tb.cell(1, 1).text = "甲", "乙"
    p.save(path)


def test_kinds_and_previews(proj):
    d = proj.materials / "汇报"
    d.mkdir(parents=True)
    for n in ("a.mp4", "b.mkv", "c.mp3", "d.docx", "e.xlsx", "f.xls", "g.pptx"):
        (d / n).write_bytes(b"\0" * 10)
    _pptx(d / "g.pptx")
    assert [files.kind_of(d / n) for n in ("a.mp4", "b.mkv", "c.mp3", "d.docx", "e.xlsx", "f.xls", "g.pptx")] == \
        ["video", "video", "audio", "docx", "xlsx", "xlsx", "pptx"]
    with TestClient(create_app(proj)) as client:
        pv = lambda n: client.get("/api/modules/汇报/preview", params={"path": n}).json()
        assert pv("a.mp4")["playable"] is True and pv("b.mkv")["playable"] is False and pv("c.mp3")["kind"] == "audio"
        x = pv("g.pptx")["pptx"]
        s1, s2 = x["slides"]
        assert (s1["title"], [t["t"] for t in s1["texts"]], s1["texts"][1]["lv"], s1["pictures"], s1["notes"]) == \
            ("第一页的标题", ["要点一", "要点二"], 1, 1, "讲稿在这")
        assert s2["title"] == "第二页有表格" and s2["tables"][0][0][0] == "甲" and s2["tables"][0][1][1] == "乙"
        img = client.get("/files/汇报/g.pptx", params={"slide": 1, "pic": 0})
        assert img.status_code == 200 and img.headers["content-type"] == "image/png" and img.content[:4] == b"\x89PNG"
        assert client.get("/files/汇报/g.pptx", params={"slide": 1, "pic": 5}).status_code == 404
        assert client.get("/files/汇报/g.pptx", params={"slide": 9}).status_code == 404
        (d / "坏的.pptx").write_bytes(b"not a pptx")
        assert "读不出来" in pv("坏的.pptx")["error"]


def test_video_can_seek(proj):
    d = proj.materials / "素材"
    d.mkdir(parents=True)
    (d / "v.mp4").write_bytes(bytes(range(256)) * 40)
    with TestClient(create_app(proj)) as client:
        r = client.get("/files/素材/v.mp4", headers={"Range": "bytes=100-199"})
        assert r.status_code == 206 and len(r.content) == 100 and r.content[0] == 100      # 拖进度条靠这个


def test_open_local_only_opens_documents(proj, monkeypatch):
    d = proj.materials / "汇报"
    d.mkdir(parents=True)
    (d / "讲稿.docx").write_bytes(b"x")
    (d / "坏.bat").write_text("echo hi", encoding="utf-8")
    opened = []
    monkeypatch.setattr(office, "open_local", lambda f: opened.append(f.name) if f.suffix.lower() in office.OPENABLE
                        else (_ for _ in ()).throw(ValueError("「.bat」这种文件不从网页打开")))
    with TestClient(create_app(proj)) as client:
        assert client.post("/api/open-local", json={"module": "汇报", "path": "讲稿.docx"}).json()["ok"] is True
        assert client.post("/api/open-local", json={"module": "汇报", "path": "坏.bat"}).status_code == 400
        assert client.post("/api/open-local", json={"path": "资料/汇报/讲稿.docx"}).json()["name"] == "讲稿.docx"
        assert client.post("/api/open-local", json={"module": "汇报", "path": "../../AGENTS.md"}).status_code >= 400
    assert opened == ["讲稿.docx", "讲稿.docx"]
    assert ".bat" not in office.OPENABLE and ".exe" not in office.OPENABLE and ".ps1" not in office.OPENABLE
